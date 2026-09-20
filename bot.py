"""
Telegram-бот учёта работ и расходов монтажников.

/start -> меню:
  📋 Отчёт о работе — объект, дата, время с/по, расходы, описание
  💰 Добавить расход — объект, дата, сумма, описание расхода

Администратор: /add, /objects, /report.
"""
import asyncio
import io
import logging
import os
import re
import sqlite3
from datetime import date, datetime, timedelta

from aiogram import Bot, Dispatcher, F, Router
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (BotCommand, BotCommandScopeChat, BufferedInputFile,
                           CallbackQuery, Message)
from aiogram.utils.keyboard import InlineKeyboardBuilder
from openpyxl import Workbook
from openpyxl.styles import Font

# ================== НАСТРОЙКИ ==================
# Бот @Asakom_report_bot. Токен задаётся переменной окружения BOT_TOKEN
# (в панели BotHost — поле «Bot Token»), чтобы не хранить его в коде.
BOT_TOKEN = os.getenv("BOT_TOKEN", "")
# Telegram ID главных администраторов через запятую.
# Если пусто — главным станет тот, кто первым нажмёт /start.
ADMIN_IDS = {
    int(x) for x in os.getenv("ADMIN_IDS", "603191958").replace(" ", "").split(",") if x
}
DB_PATH = os.getenv("DB_PATH", "bot.db")
# Часы, которые показываются кнопками при выборе времени
HOUR_BUTTONS = range(6, 23)
# --- Яндекс.Диск (необязательно) ---
# Токен с правами на запись: https://yandex.ru/dev/disk/poligon/ — кнопка «Получить OAuth-токен».
# Задаётся переменной окружения YADISK_TOKEN. Без него бот работает как обычно.
YADISK_TOKEN = os.getenv("YADISK_TOKEN", "")
YADISK_FOLDER = os.getenv("YADISK_FOLDER", "disk:/Асаком отчёты")
YADISK_FILE = os.getenv("YADISK_FILE", "Отчёт монтажников.xlsx")
# Как часто проверять, не пора ли обновить файл на Диске (секунды)
SYNC_SECONDS = int(os.getenv("SYNC_SECONDS", "60"))
# ===============================================

router = Router()

# Команды в синей кнопке «Меню». Монтажники видят только первые три.
USER_COMMANDS = [
    BotCommand(command="start", description="Меню"),
    BotCommand(command="name", description="Изменить ФИО"),
    BotCommand(command="cancel", description="Отменить ввод"),
]
ADMIN_COMMANDS = USER_COMMANDS + [
    BotCommand(command="admin", description="Панель администратора"),
    BotCommand(command="disk", description="Обновить файл на Яндекс.Диске"),
]


async def apply_commands(bot, user_id):
    """Показывает человеку тот набор команд, который ему положен."""
    try:
        await bot.set_my_commands(
            ADMIN_COMMANDS if is_admin(user_id) else USER_COMMANDS,
            scope=BotCommandScopeChat(chat_id=user_id),
        )
    except Exception:
        pass


# ---------- База данных ----------
def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as c:
        c.execute(
            """CREATE TABLE IF NOT EXISTS objects (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1)"""
        )
        c.execute(
            """CREATE TABLE IF NOT EXISTS admins (
                user_id INTEGER PRIMARY KEY,
                fio TEXT,
                added_by TEXT,
                created_at TEXT)"""
        )
        c.execute(
            """CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                fio TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'pending',
                created_at TEXT)"""
        )
        if "status" not in [r[1] for r in c.execute("PRAGMA table_info(users)")]:
            # в базе от прежней версии все уже работавшие считаются одобренными
            c.execute("ALTER TABLE users ADD COLUMN status TEXT NOT NULL DEFAULT 'approved'")
        c.execute(
            """CREATE TABLE IF NOT EXISTS reports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER, user_name TEXT,
                object_id INTEGER, work_date TEXT,
                time_from TEXT, time_to TEXT, hours REAL,
                expenses REAL, exp_note TEXT, description TEXT,
                entered_by TEXT, created_at TEXT)"""
        )
        # для баз, созданных прежней версией бота
        cols = [r[1] for r in c.execute("PRAGMA table_info(reports)")]
        if "exp_note" not in cols:
            c.execute("ALTER TABLE reports ADD COLUMN exp_note TEXT")
        if "entered_by" not in cols:
            c.execute("ALTER TABLE reports ADD COLUMN entered_by TEXT")
        c.execute(
            """CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER, user_name TEXT,
                object_id INTEGER, exp_date TEXT,
                amount REAL, category TEXT, comment TEXT,
                photo_id TEXT, entered_by TEXT, created_at TEXT)"""
        )
        if "entered_by" not in [r[1] for r in c.execute("PRAGMA table_info(expenses)")]:
            c.execute("ALTER TABLE expenses ADD COLUMN entered_by TEXT")


def active_objects():
    with db() as c:
        return c.execute("SELECT id, name FROM objects WHERE active=1 ORDER BY name").fetchall()


def object_name(obj_id):
    with db() as c:
        row = c.execute("SELECT name FROM objects WHERE id=?", (obj_id,)).fetchone()
        return row["name"] if row else "?"


def owner_ids():
    """Главные администраторы: из настроек или тот, кто первым запустил бота."""
    if ADMIN_IDS:
        return set(ADMIN_IDS)
    with db() as c:
        return {r["user_id"] for r in c.execute(
            "SELECT user_id FROM admins WHERE added_by='владелец'")}


def is_owner(user_id):
    """Главного администратора снять через бот нельзя."""
    return user_id in owner_ids()


def claim_owner(user_id, fio):
    """Первый, кто открыл бота, становится главным администратором."""
    with db() as c:
        if c.execute("SELECT 1 FROM admins LIMIT 1").fetchone():
            return False
        c.execute(
            "INSERT INTO admins (user_id, fio, added_by, created_at) VALUES (?,?,?,?)",
            (user_id, fio, "владелец", datetime.now().strftime("%Y-%m-%d %H:%M")),
        )
    return True


def is_admin(user_id):
    """Начальник: либо прописан в настройках, либо назначен через бот."""
    if user_id in ADMIN_IDS:
        return True
    with db() as c:
        return c.execute("SELECT 1 FROM admins WHERE user_id=?", (user_id,)).fetchone() is not None


def admin_ids():
    """Все, кому уходят уведомления о новых отчётах."""
    ids = set(ADMIN_IDS)
    with db() as c:
        ids.update(r["user_id"] for r in c.execute("SELECT user_id FROM admins"))
    return ids


def get_fio(user_id):
    """ФИО монтажника из базы или None, если он ещё не зарегистрирован."""
    with db() as c:
        row = c.execute("SELECT fio FROM users WHERE user_id=?", (user_id,)).fetchone()
        return row["fio"] if row else None


def save_fio(user_id, fio):
    """Сохраняет ФИО. Начальники одобряются сразу, остальные ждут подтверждения."""
    status = "approved" if is_admin(user_id) else "pending"
    with db() as c:
        c.execute(
            "INSERT INTO users (user_id, fio, status, created_at) VALUES (?,?,?,?) "
            "ON CONFLICT(user_id) DO UPDATE SET fio=excluded.fio",
            (user_id, fio, status, datetime.now().strftime("%Y-%m-%d %H:%M")),
        )


def user_status(user_id):
    """'approved' — допущен, 'pending' — ждёт решения, 'rejected' — отказано, None — не заходил."""
    if is_admin(user_id):
        return "approved"
    with db() as c:
        row = c.execute("SELECT status FROM users WHERE user_id=?", (user_id,)).fetchone()
        return row["status"] if row else None


def set_status(user_id, status):
    with db() as c:
        c.execute("UPDATE users SET status=? WHERE user_id=?", (status, user_id))


def pending_users():
    with db() as c:
        return c.execute(
            "SELECT user_id, fio, created_at FROM users WHERE status='pending' ORDER BY created_at"
        ).fetchall()


# ---------- Состояния ----------
class Reg(StatesGroup):           # первый вход
    fio = State()


class Form(StatesGroup):          # отчёт о работе
    obj = State()
    work_date = State()
    time_from = State()
    time_to = State()
    expenses = State()
    exp_note = State()
    description = State()
    confirm = State()


class Exp(StatesGroup):           # отдельный расход
    obj = State()
    exp_date = State()
    amount = State()
    comment = State()
    confirm = State()


# ---------- Вспомогательное ----------
def parse_time(text):
    m = re.fullmatch(r"\s*(\d{1,2})[:.](\d{2})\s*", text or "")
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    if 0 <= h <= 23 and 0 <= mi <= 59:
        return f"{h:02d}:{mi:02d}"
    return None


def parse_amount(text):
    try:
        v = float((text or "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None
    return v if v >= 0 else None


def parse_date(text):
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%d.%m"):
        try:
            d = datetime.strptime((text or "").strip(), fmt).date()
        except ValueError:
            continue
        if fmt == "%d.%m":
            d = d.replace(year=date.today().year)
        return d
    return None


def minutes(t):
    h, m = map(int, t.split(":"))
    return h * 60 + m


def time_kb(prefix, after=None):
    kb = InlineKeyboardBuilder()
    for h in HOUR_BUTTONS:
        t = f"{h:02d}:00"
        if after is None or minutes(t) > minutes(after):
            kb.button(text=t, callback_data=f"{prefix}:{t}")
    kb.adjust(4)
    return kb.as_markup()


def date_kb(prefix):
    kb = InlineKeyboardBuilder()
    today = date.today()
    for label, delta in (("Сегодня", 0), ("Вчера", 1), ("Позавчера", 2)):
        d = today - timedelta(days=delta)
        kb.button(text=f"{label} ({d:%d.%m})", callback_data=f"{prefix}:{d.isoformat()}")
    kb.adjust(1)
    return kb.as_markup()


def objects_kb(prefix):
    kb = InlineKeyboardBuilder()
    for o in active_objects():
        kb.button(text=o["name"], callback_data=f"{prefix}:{o['id']}")
    kb.adjust(1)
    return kb.as_markup()


def fmt_date(iso):
    return datetime.strptime(iso, "%Y-%m-%d").strftime("%d.%m.%Y")


def confirm_kb():
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Сохранить", callback_data="ok")
    kb.button(text="❌ Отмена", callback_data="no")
    return kb.as_markup()


def work_summary(d):
    return (
        f"🏗 Объект: {object_name(d['obj'])}\n"
        f"📅 Дата: {fmt_date(d['work_date'])}\n"
        f"🕐 Время: {d['time_from']}–{d['time_to']} ({d['hours']:g} ч)\n"
        f"💰 Расходы: {d['expenses']:g}"
        + (f" — {d['exp_note']}" if d.get("exp_note") and d["exp_note"] != "—" else "")
        + f"\n📝 Описание и с кем работал: {d['description']}"
    )


def exp_summary(d):
    return (
        f"🏗 Объект: {object_name(d['obj'])}\n"
        f"📅 Дата: {fmt_date(d['exp_date'])}\n"
        f"💰 Сумма: {d['amount']:g}\n"
        f"📝 На что: {d['comment']}"
    )


# ---------- Меню ----------
def menu_kb(user_id=None):
    kb = InlineKeyboardBuilder()
    kb.button(text="📋 Отчёт о работе", callback_data="menu:work")
    kb.button(text="💰 Добавить расход", callback_data="menu:exp")
    kb.button(text="🗂 Мои записи", callback_data="menu:my")
    if user_id is not None and is_admin(user_id):
        kb.button(text="✍️ Внести за монтажника", callback_data="menu:forother")
    kb.adjust(1)
    return kb.as_markup()


@router.message(Command("cancel"))
async def cancel(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("Отменено. Что дальше?", reply_markup=menu_kb(message.from_user.id))


@router.message(CommandStart())
async def start(message: Message, state: FSMContext, bot: Bot):
    await state.clear()
    await apply_commands(bot, message.from_user.id)
    status = user_status(message.from_user.id)
    if status is None:
        await state.set_state(Reg.fio)
        await message.answer(
            "Здравствуйте! Вы здесь впервые.\n\n"
            "Напишите свою фамилию, имя и отчество — они будут указываться в отчётах.\n"
            "Например: Иванов Иван Иванович"
        )
        return
    if status != "approved":
        await message.answer(WAIT_TEXT if status == "pending" else DENIED_TEXT)
        return
    await show_menu(message, greet=True)


@router.message(Reg.fio, F.text)
async def reg_fio(message: Message, state: FSMContext, bot: Bot):
    fio = " ".join((message.text or "").split())
    if len(fio) < 5 or len(fio.split()) < 2:
        await message.answer("Напишите полностью, например: Иванов Иван Иванович")
        return
    known = user_status(message.from_user.id) is not None
    # самый первый человек в пустой базе становится главным администратором
    if not ADMIN_IDS and claim_owner(message.from_user.id, fio):
        await message.answer(
            f"{fio}, вы первый в этом боте, поэтому назначены главным администратором.\n"
            "Управление объектами, отчётами и доступами — команда /admin"
        )
    save_fio(message.from_user.id, fio)
    await apply_commands(bot, message.from_user.id)
    await state.clear()
    status = user_status(message.from_user.id)

    if status == "approved":
        await message.answer(f"Записал: {fio}\nИзменить можно командой /name")
        await show_menu(message)
        return
    if status == "rejected":
        await message.answer(DENIED_TEXT)
        return

    # новый человек — отправляем заявку начальникам
    await message.answer(f"Записал: {fio}\n\n" + WAIT_TEXT)
    if known:
        return
    username = f"\nTelegram: @{message.from_user.username}" if message.from_user.username else ""
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Принять", callback_data=f"u:ok:{message.from_user.id}")
    kb.button(text="⛔️ Отклонить", callback_data=f"u:no:{message.from_user.id}")
    kb.adjust(2)
    for admin in admin_ids():
        if admin == message.from_user.id:
            continue
        try:
            await bot.send_message(
                admin,
                f"🔔 Новая заявка на доступ\n\n👤 {fio}{username}\n\n"
                "Допустить этого человека к внесению отчётов?",
                reply_markup=kb.as_markup(),
            )
        except Exception:
            pass


WAIT_TEXT = (
    "Заявка отправлена руководителю. Как только её подтвердят, "
    "бот пришлёт сообщение, и можно будет вносить отчёты."
)
DENIED_TEXT = "Доступ к боту не подтверждён. Обратитесь к своему руководителю."


@router.message(Command("name"))
async def change_name(message: Message, state: FSMContext):
    if user_status(message.from_user.id) == "rejected":
        await message.answer(DENIED_TEXT)
        return
    await state.set_state(Reg.fio)
    current = get_fio(message.from_user.id)
    await message.answer(
        (f"Сейчас записано: {current}\n\n" if current else "")
        + "Напишите ФИО заново:"
    )


async def show_menu(message: Message, greet: bool = False):
    text = ""
    if greet:
        text = f"{get_fio(message.from_user.id)}\n\n"
    admin = is_admin(message.from_user.id)
    if admin:
        text += "Вы администратор. Управление объектами и отчётами — /admin\n\n"
    if not active_objects():
        await message.answer(
            text + ("Объектов пока нет — добавьте первый." if admin
                    else "Пока нет ни одного объекта, обратитесь к руководителю."),
            reply_markup=admin_kb() if admin else None,
        )
        return
    await message.answer(text + "Что хотите сделать?", reply_markup=menu_kb(message.from_user.id))


@router.callback_query(F.data == "menu:work")
async def menu_work(cb: CallbackQuery, state: FSMContext):
    await state.clear()
    status = user_status(cb.from_user.id)
    if status is None:
        await state.set_state(Reg.fio)
        await cb.message.answer("Сначала напишите своё ФИО, например: Иванов Иван Иванович")
        await cb.answer()
        return
    if status != "approved":
        await cb.message.answer(WAIT_TEXT if status == "pending" else DENIED_TEXT)
        await cb.answer()
        return
    await state.set_state(Form.obj)
    await cb.message.edit_text("📋 Отчёт о работе\n\nВыберите объект:",
                               reply_markup=objects_kb("obj"))
    await cb.answer()


@router.callback_query(F.data == "menu:exp")
async def menu_exp(cb: CallbackQuery, state: FSMContext):
    await state.clear()
    status = user_status(cb.from_user.id)
    if status is None:
        await state.set_state(Reg.fio)
        await cb.message.answer("Сначала напишите своё ФИО, например: Иванов Иван Иванович")
        await cb.answer()
        return
    if status != "approved":
        await cb.message.answer(WAIT_TEXT if status == "pending" else DENIED_TEXT)
        await cb.answer()
        return
    await state.set_state(Exp.obj)
    await cb.message.edit_text("💰 Добавить расход\n\nВыберите объект:",
                               reply_markup=objects_kb("eobj"))
    await cb.answer()


# ---------- Начальник вносит запись за монтажника ----------
def approved_users():
    with db() as c:
        return c.execute(
            "SELECT user_id, fio FROM users WHERE status='approved' ORDER BY fio").fetchall()


@router.callback_query(F.data == "menu:forother")
async def menu_for_other(cb: CallbackQuery, state: FSMContext):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    await state.clear()
    people = approved_users()
    if not people:
        await cb.message.edit_text("Нет ни одного допущенного монтажника.",
                                   reply_markup=menu_kb(cb.from_user.id))
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    for r in people:
        kb.button(text=r["fio"], callback_data=f"fo:{r['user_id']}")
    kb.adjust(1)
    await cb.message.edit_text(
        "✍️ Внести запись за монтажника\n\nЗа кого вносим?", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.startswith("fo:"))
async def for_other_picked(cb: CallbackQuery, state: FSMContext):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    uid = int(cb.data.split(":")[1])
    await state.update_data(for_user=uid, for_fio=get_fio(uid) or f"ID {uid}")
    kb = InlineKeyboardBuilder()
    kb.button(text="📋 Отчёт о работе", callback_data="fo_kind:work")
    kb.button(text="💰 Расход", callback_data="fo_kind:exp")
    kb.adjust(1)
    await cb.message.edit_text(
        f"Вносим за: {get_fio(uid)}\n\nЧто вносим?", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.startswith("fo_kind:"))
async def for_other_kind(cb: CallbackQuery, state: FSMContext):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    data = await state.get_data()
    if cb.data.endswith("work"):
        await state.set_state(Form.obj)
        await cb.message.edit_text(
            f"📋 Отчёт за {data.get('for_fio')}\n\nВыберите объект:",
            reply_markup=objects_kb("obj"))
    else:
        await state.set_state(Exp.obj)
        await cb.message.edit_text(
            f"💰 Расход за {data.get('for_fio')}\n\nВыберите объект:",
            reply_markup=objects_kb("eobj"))
    await cb.answer()


# ---------- Отчёт о работе ----------
@router.callback_query(Form.obj, F.data.startswith("obj:"))
async def choose_obj(cb: CallbackQuery, state: FSMContext):
    obj_id = int(cb.data.split(":")[1])
    await state.update_data(obj=obj_id)
    await state.set_state(Form.work_date)
    await cb.message.edit_text(
        f"Объект: {object_name(obj_id)}\n\nВыберите дату или напишите её, например 17.09.2026:",
        reply_markup=date_kb("date"),
    )
    await cb.answer()


async def ask_time_from(target: Message, state: FSMContext, iso):
    await state.update_data(work_date=iso)
    await state.set_state(Form.time_from)
    await target.answer(
        f"Дата: {fmt_date(iso)}\n\nС какого времени работали? Выберите или напишите, например 8:30",
        reply_markup=time_kb("tf"),
    )


@router.callback_query(Form.work_date, F.data.startswith("date:"))
async def date_btn(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_reply_markup()
    await ask_time_from(cb.message, state, cb.data.split(":")[1])
    await cb.answer()


@router.message(Form.work_date)
async def date_text(message: Message, state: FSMContext):
    d = parse_date(message.text)
    if not d:
        await message.answer("Не понял дату. Формат: 17.09.2026")
        return
    if d > date.today():
        await message.answer("Дата не может быть в будущем.")
        return
    await ask_time_from(message, state, d.isoformat())


async def ask_time_to(target: Message, state: FSMContext, t):
    await state.update_data(time_from=t)
    await state.set_state(Form.time_to)
    await target.answer(
        f"Начало: {t}\n\nДо какого времени работали? Выберите или напишите, например 17:30",
        reply_markup=time_kb("tt", after=t),
    )


@router.callback_query(Form.time_from, F.data.startswith("tf:"))
async def tf_btn(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_reply_markup()
    await ask_time_to(cb.message, state, cb.data[3:])
    await cb.answer()


@router.message(Form.time_from)
async def tf_text(message: Message, state: FSMContext):
    t = parse_time(message.text)
    if not t:
        await message.answer("Формат времени: 8:30 или 08:30")
        return
    await ask_time_to(message, state, t)


async def ask_work_expenses(target: Message, state: FSMContext, t):
    data = await state.get_data()
    diff = minutes(t) - minutes(data["time_from"])
    if diff <= 0:
        await target.answer("Время окончания должно быть позже начала. Введите ещё раз:")
        return
    await state.update_data(time_to=t, hours=round(diff / 60, 2))
    await state.set_state(Form.expenses)
    kb = InlineKeyboardBuilder()
    kb.button(text="Без расходов", callback_data="exp0")
    await target.answer(
        f"Время: {data['time_from']}–{t} ({diff / 60:g} ч)\n\n"
        "Сумма расходов за этот день? Напишите число, например 1500",
        reply_markup=kb.as_markup(),
    )


@router.callback_query(Form.time_to, F.data.startswith("tt:"))
async def tt_btn(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_reply_markup()
    await ask_work_expenses(cb.message, state, cb.data[3:])
    await cb.answer()


@router.message(Form.time_to)
async def tt_text(message: Message, state: FSMContext):
    t = parse_time(message.text)
    if not t:
        await message.answer("Формат времени: 17:30")
        return
    await ask_work_expenses(message, state, t)


async def after_amount(target: Message, state: FSMContext, amount):
    """Если расходы есть — спрашиваем, на что; если нет — сразу описание работы."""
    await state.update_data(expenses=amount)
    if amount > 0:
        await state.set_state(Form.exp_note)
        await target.answer(
            f"Расходы: {amount:g}\n\nНа что потрачены эти деньги? "
            "Например: анкеры и крепёж, заправка машины"
        )
    else:
        await state.update_data(exp_note="—")
        await state.set_state(Form.description)
        await target.answer(
            "Расходов нет.\n\nОпишите выполненную работу и укажите, "
            "с кем работали (если были напарники):"
        )


@router.callback_query(Form.expenses, F.data == "exp0")
async def exp_zero(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_reply_markup()
    await after_amount(cb.message, state, 0.0)
    await cb.answer()


@router.message(Form.expenses)
async def exp_text(message: Message, state: FSMContext):
    amount = parse_amount(message.text)
    if amount is None:
        await message.answer("Введите сумму числом, например 1500 или 0")
        return
    await after_amount(message, state, amount)


@router.message(Form.exp_note, F.text)
async def work_exp_note(message: Message, state: FSMContext):
    await state.update_data(exp_note=message.text.strip())
    await state.set_state(Form.description)
    await message.answer(
        "Принято.\n\nТеперь опишите выполненную работу и укажите, "
        "с кем работали (если были напарники):"
    )


@router.message(Form.description, F.text)
async def description(message: Message, state: FSMContext):
    await state.update_data(description=message.text.strip())
    await state.set_state(Form.confirm)
    data = await state.get_data()
    await message.answer("Проверьте:\n\n" + work_summary(data), reply_markup=confirm_kb())


@router.callback_query(Form.confirm, F.data.in_({"ok", "no"}))
async def work_confirm(cb: CallbackQuery, state: FSMContext, bot: Bot):
    if cb.data == "no":
        await state.clear()
        await cb.message.edit_text("Отменено.")
        await cb.message.answer("Что дальше?", reply_markup=menu_kb(cb.from_user.id))
        await cb.answer()
        return
    d = await state.get_data()
    author = cb.from_user
    author_fio = get_fio(author.id) or author.full_name
    # начальник мог вносить запись за монтажника
    uid = d.get("for_user", author.id)
    fio = d.get("for_fio") or author_fio
    entered_by = author_fio if uid != author.id else None
    with db() as c:
        c.execute(
            """INSERT INTO reports (user_id, user_name, object_id, work_date, time_from,
               time_to, hours, expenses, exp_note, description, entered_by, created_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (uid, fio, d["obj"], d["work_date"], d["time_from"],
             d["time_to"], d["hours"], d["expenses"], d.get("exp_note", "—"),
             d["description"], entered_by, datetime.now().strftime("%Y-%m-%d %H:%M")),
        )
    mark_changed()
    await state.clear()
    note = f"\n\nЗапись внесена за {fio}." if entered_by else ""
    await cb.message.edit_text("✅ Отчёт сохранён!\n\n" + work_summary(d) + note)
    await cb.message.answer("Что дальше?", reply_markup=menu_kb(cb.from_user.id))
    await cb.answer()
    tail = f"\n\n(внёс {entered_by})" if entered_by else ""
    for admin in admin_ids():
        if admin != author.id:
            try:
                await bot.send_message(
                    admin, f"📋 Новый отчёт — {fio}:\n\n" + work_summary(d) + tail
                )
            except Exception:
                pass
    if entered_by and uid != author.id:
        try:
            await bot.send_message(
                uid, f"{entered_by} внёс за вас отчёт:\n\n" + work_summary(d)
            )
        except Exception:
            pass


# ---------- Отдельный расход ----------
@router.callback_query(Exp.obj, F.data.startswith("eobj:"))
async def exp_obj(cb: CallbackQuery, state: FSMContext):
    obj_id = int(cb.data.split(":")[1])
    await state.update_data(obj=obj_id)
    await state.set_state(Exp.exp_date)
    await cb.message.edit_text(
        f"Объект: {object_name(obj_id)}\n\n"
        "Дата расхода — выберите или напишите, например 17.09.2026:",
        reply_markup=date_kb("edate"),
    )
    await cb.answer()


async def ask_exp_amount(target: Message, state: FSMContext, iso):
    await state.update_data(exp_date=iso)
    await state.set_state(Exp.amount)
    await target.answer(f"Дата: {fmt_date(iso)}\n\nСумма расхода? Напишите число, например 2450")


@router.callback_query(Exp.exp_date, F.data.startswith("edate:"))
async def exp_date_btn(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_reply_markup()
    await ask_exp_amount(cb.message, state, cb.data.split(":")[1])
    await cb.answer()


@router.message(Exp.exp_date)
async def exp_date_text(message: Message, state: FSMContext):
    d = parse_date(message.text)
    if not d:
        await message.answer("Не понял дату. Формат: 17.09.2026")
        return
    if d > date.today():
        await message.answer("Дата не может быть в будущем.")
        return
    await ask_exp_amount(message, state, d.isoformat())


@router.message(Exp.amount)
async def exp_amount(message: Message, state: FSMContext):
    amount = parse_amount(message.text)
    if not amount:
        await message.answer("Введите сумму числом больше нуля, например 2450")
        return
    await state.update_data(amount=amount)
    await state.set_state(Exp.comment)
    await message.answer(
        f"Сумма: {amount:g}\n\n"
        "Опишите расход — на что потрачено."
    )


async def exp_show_confirm(target: Message, state: FSMContext):
    await state.set_state(Exp.confirm)
    data = await state.get_data()
    await target.answer("Проверьте:\n\n" + exp_summary(data), reply_markup=confirm_kb())


@router.message(Exp.comment, F.text)
async def exp_comment(message: Message, state: FSMContext):
    await state.update_data(comment=message.text.strip())
    await exp_show_confirm(message, state)


@router.callback_query(Exp.confirm, F.data.in_({"ok", "no"}))
async def exp_confirm(cb: CallbackQuery, state: FSMContext, bot: Bot):
    if cb.data == "no":
        await state.clear()
        await cb.message.edit_text("Отменено.")
        await cb.message.answer("Что дальше?", reply_markup=menu_kb(cb.from_user.id))
        await cb.answer()
        return
    d = await state.get_data()
    author = cb.from_user
    author_fio = get_fio(author.id) or author.full_name
    uid = d.get("for_user", author.id)
    fio = d.get("for_fio") or author_fio
    entered_by = author_fio if uid != author.id else None
    with db() as c:
        c.execute(
            """INSERT INTO expenses (user_id, user_name, object_id, exp_date, amount,
               category, comment, photo_id, entered_by, created_at)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (uid, fio, d["obj"], d["exp_date"], d["amount"],
             "", d["comment"], None, entered_by,
             datetime.now().strftime("%Y-%m-%d %H:%M")),
        )
    mark_changed()
    await state.clear()
    note = f"\n\nЗапись внесена за {fio}." if entered_by else ""
    await cb.message.edit_text("✅ Расход сохранён!\n\n" + exp_summary(d) + note)
    await cb.message.answer("Что дальше?", reply_markup=menu_kb(cb.from_user.id))
    await cb.answer()
    tail = f"\n\n(внёс {entered_by})" if entered_by else ""
    for admin in admin_ids():
        if admin == author.id:
            continue
        try:
            await bot.send_message(admin, f"💰 Новый расход — {fio}:\n\n" + exp_summary(d) + tail)
        except Exception:
            pass
    if entered_by and uid != author.id:
        try:
            await bot.send_message(uid, f"{entered_by} внёс за вас расход:\n\n" + exp_summary(d))
        except Exception:
            pass


# ---------- Мои записи: посмотреть, исправить, удалить ----------
def my_records(user_id, limit=10):
    """Последние записи человека: работы и расходы вперемешку, новые сверху."""
    with db() as c:
        works = c.execute(
            """SELECT r.id, r.work_date AS d, r.hours, r.expenses, r.description,
                      r.time_from, r.time_to, o.name AS obj
               FROM reports r LEFT JOIN objects o ON o.id=r.object_id
               WHERE r.user_id=? ORDER BY r.id DESC LIMIT ?""", (user_id, limit)).fetchall()
        exps = c.execute(
            """SELECT e.id, e.exp_date AS d, e.amount, e.comment, o.name AS obj
               FROM expenses e LEFT JOIN objects o ON o.id=e.object_id
               WHERE e.user_id=? ORDER BY e.id DESC LIMIT ?""", (user_id, limit)).fetchall()
    items = [("w", r) for r in works] + [("e", r) for r in exps]
    items.sort(key=lambda t: (t[1]["d"], t[1]["id"]), reverse=True)
    return items[:limit]


def record_line(kind, r):
    if kind == "w":
        return (f"📋 {fmt_date(r['d'])} · {r['obj']} · {r['hours']:g} ч"
                + (f" · {r['expenses']:g} ₽" if r["expenses"] else ""))
    return f"💰 {fmt_date(r['d'])} · {r['obj']} · {r['amount']:g} ₽"


async def show_my_records(target: Message, user_id, edit=False):
    items = my_records(user_id)
    if not items:
        text = "У вас пока нет записей."
        await (target.edit_text(text, reply_markup=menu_kb(user_id)) if edit
               else target.answer(text, reply_markup=menu_kb(user_id)))
        return
    kb = InlineKeyboardBuilder()
    for kind, r in items:
        kb.button(text=record_line(kind, r), callback_data=f"rec:{kind}:{r['id']}")
    kb.adjust(1)
    text = "🗂 Ваши последние записи.\n\nНажмите на запись, чтобы исправить или удалить её:"
    await (target.edit_text(text, reply_markup=kb.as_markup()) if edit
           else target.answer(text, reply_markup=kb.as_markup()))


@router.callback_query(F.data == "menu:my")
async def menu_my(cb: CallbackQuery, state: FSMContext):
    await state.clear()
    if user_status(cb.from_user.id) != "approved":
        await cb.answer()
        return
    await show_my_records(cb.message, cb.from_user.id, edit=True)
    await cb.answer()


def load_record(kind, rec_id):
    with db() as c:
        if kind == "w":
            return c.execute(
                """SELECT r.*, o.name AS obj FROM reports r
                   LEFT JOIN objects o ON o.id=r.object_id WHERE r.id=?""", (rec_id,)).fetchone()
        return c.execute(
            """SELECT e.*, o.name AS obj FROM expenses e
               LEFT JOIN objects o ON o.id=e.object_id WHERE e.id=?""", (rec_id,)).fetchone()


@router.callback_query(F.data.startswith("rec:"))
async def record_card(cb: CallbackQuery):
    _, kind, rec_id = cb.data.split(":")
    r = load_record(kind, int(rec_id))
    if not r or (r["user_id"] != cb.from_user.id and not is_admin(cb.from_user.id)):
        await cb.answer("Запись не найдена", show_alert=True)
        return
    if kind == "w":
        text = (f"📋 Отчёт о работе\n\n🏗 {r['obj']}\n📅 {fmt_date(r['work_date'])}\n"
                f"🕐 {r['time_from']}–{r['time_to']} ({r['hours']:g} ч)\n"
                f"💰 Расходы: {r['expenses']:g}"
                + (f" — {r['exp_note']}" if r["exp_note"] and r["exp_note"] != "—" else "")
                + f"\n📝 {r['description']}")
    else:
        text = (f"💰 Расход\n\n🏗 {r['obj']}\n📅 {fmt_date(r['exp_date'])}\n"
                f"Сумма: {r['amount']:g}\n📝 {r['comment']}")
    kb = InlineKeyboardBuilder()
    kb.button(text="✏️ Исправить (внести заново)", callback_data=f"recfix:{kind}:{rec_id}")
    kb.button(text="🗑 Удалить", callback_data=f"recdel:{kind}:{rec_id}")
    kb.button(text="⬅️ К списку", callback_data="menu:my")
    kb.adjust(1)
    await cb.message.edit_text(text, reply_markup=kb.as_markup())
    await cb.answer()


async def drop_record(kind, rec_id):
    with db() as c:
        c.execute(f"DELETE FROM {'reports' if kind == 'w' else 'expenses'} WHERE id=?", (rec_id,))
    mark_changed()


@router.callback_query(F.data.startswith("recdel:"))
async def record_delete(cb: CallbackQuery, bot: Bot):
    _, kind, rec_id = cb.data.split(":")
    r = load_record(kind, int(rec_id))
    if not r or (r["user_id"] != cb.from_user.id and not is_admin(cb.from_user.id)):
        await cb.answer("Запись не найдена", show_alert=True)
        return
    await drop_record(kind, int(rec_id))
    who = get_fio(cb.from_user.id) or cb.from_user.full_name
    await cb.message.edit_text("🗑 Запись удалена.")
    await show_my_records(cb.message, cb.from_user.id)
    await cb.answer()
    line = record_line(kind, {**dict(r), "obj": r["obj"],
                              "d": r["work_date"] if kind == "w" else r["exp_date"]})
    for admin in admin_ids():
        if admin != cb.from_user.id:
            try:
                await bot.send_message(admin, f"🗑 {who} удалил запись:\n{line}")
            except Exception:
                pass


@router.callback_query(F.data.startswith("recfix:"))
async def record_fix(cb: CallbackQuery, state: FSMContext, bot: Bot):
    _, kind, rec_id = cb.data.split(":")
    r = load_record(kind, int(rec_id))
    if not r or (r["user_id"] != cb.from_user.id and not is_admin(cb.from_user.id)):
        await cb.answer("Запись не найдена", show_alert=True)
        return
    await drop_record(kind, int(rec_id))
    await state.clear()
    # если начальник правит чужую запись — новая сохранится на того же человека
    if r["user_id"] != cb.from_user.id:
        await state.update_data(for_user=r["user_id"], for_fio=r["user_name"])
    if kind == "w":
        await state.set_state(Form.obj)
        await cb.message.edit_text(
            "Старая запись удалена, вносим заново.\n\nВыберите объект:",
            reply_markup=objects_kb("obj"))
    else:
        await state.set_state(Exp.obj)
        await cb.message.edit_text(
            "Старый расход удалён, вносим заново.\n\nВыберите объект:",
            reply_markup=objects_kb("eobj"))
    await cb.answer()


# ---------- Администратор ----------
class Adm(StatesGroup):
    add_obj = State()
    rename = State()


def admin_kb():
    kb = InlineKeyboardBuilder()
    waiting = len(pending_users())
    if waiting:
        kb.button(text=f"🔔 Заявки на доступ ({waiting})", callback_data="a:req")
    kb.button(text="➕ Добавить объект", callback_data="a:new")
    kb.button(text="🏗 Объекты в работе", callback_data="a:list")
    kb.button(text="📦 Архив объектов", callback_data="a:arch")
    kb.button(text="📊 Отчёт в Excel", callback_data="a:rep")
    kb.button(text="👥 Монтажники", callback_data="a:people")
    kb.button(text="👑 Начальники", callback_data="a:bosses")
    kb.adjust(1)
    return kb.as_markup()


ADMIN_HELP = (
    "Панель администратора.\n\n"
    "Быстрые команды:\n"
    "/add Название — добавить объект одной строкой\n"
    "/objects — список объектов\n"
    "/report — выгрузка в Excel, /report 09.2026 — за месяц\n"
    "/admin — эта панель"
)


@router.message(Command("admin"))
async def admin_panel(message: Message, state: FSMContext):
    if not is_admin(message.from_user.id):
        return
    await state.clear()
    await message.answer(ADMIN_HELP, reply_markup=admin_kb())


async def back_to_admin(target: Message):
    await target.answer("Что дальше?", reply_markup=admin_kb())


def save_objects(text):
    """Добавляет объекты: каждая строка — отдельный объект. Возвращает (добавленные, дубли)."""
    added, dupes = [], []
    existing = {o["name"].lower() for o in active_objects()}
    for line in (text or "").splitlines():
        name = " ".join(line.split())
        if not name:
            continue
        if name.lower() in existing:
            dupes.append(name)
            continue
        with db() as c:
            c.execute("INSERT INTO objects (name) VALUES (?)", (name,))
        existing.add(name.lower())
        added.append(name)
    return added, dupes


@router.message(Command("add"))
async def add_object(message: Message, state: FSMContext):
    if not is_admin(message.from_user.id):
        return
    name = (message.text or "").partition(" ")[2].strip()
    if not name:
        await state.set_state(Adm.add_obj)
        await message.answer(
            "Напишите название объекта.\n"
            "Можно сразу несколько — каждый с новой строки:\n\n"
            "ЖК Солнечный, кв. 12\nОфис на Ленина, 40\nСклад Северный"
        )
        return
    added, dupes = save_objects(name)
    await message.answer(report_added(added, dupes))


def report_added(added, dupes):
    parts = []
    if added:
        parts.append("Добавлено:\n" + "\n".join(f"• {n}" for n in added))
    if dupes:
        parts.append("Уже есть в списке:\n" + "\n".join(f"• {n}" for n in dupes))
    return "\n\n".join(parts) or "Ничего не добавлено."


@router.callback_query(F.data == "a:new")
async def a_new(cb: CallbackQuery, state: FSMContext):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    await state.set_state(Adm.add_obj)
    await cb.message.edit_text(
        "Напишите название объекта.\n"
        "Можно сразу несколько — каждый с новой строки:\n\n"
        "ЖК Солнечный, кв. 12\nОфис на Ленина, 40\nСклад Северный"
    )
    await cb.answer()


@router.message(Adm.add_obj, F.text)
async def a_new_save(message: Message, state: FSMContext):
    added, dupes = save_objects(message.text)
    await state.clear()
    await message.answer(report_added(added, dupes))
    await back_to_admin(message)


@router.message(Command("objects"))
async def list_objects(message: Message):
    if not is_admin(message.from_user.id):
        return
    objs = active_objects()
    if not objs:
        await message.answer("Объектов пока нет.", reply_markup=admin_kb())
        return
    kb = InlineKeyboardBuilder()
    for o in objs:
        kb.button(text=o["name"], callback_data=f"a:obj:{o['id']}")
    kb.adjust(1)
    await message.answer("Нажмите на объект, чтобы изменить его:", reply_markup=kb.as_markup())


@router.callback_query(F.data == "a:list")
async def a_list(cb: CallbackQuery, state: FSMContext):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    await state.clear()
    objs = active_objects()
    if not objs:
        await cb.message.edit_text("Объектов пока нет.", reply_markup=admin_kb())
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    for o in objs:
        kb.button(text=o["name"], callback_data=f"a:obj:{o['id']}")
    kb.button(text="⬅️ Назад", callback_data="a:back")
    kb.adjust(1)
    await cb.message.edit_text("Нажмите на объект, чтобы изменить его:", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data == "a:back")
async def a_back(cb: CallbackQuery, state: FSMContext):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    await state.clear()
    await cb.message.edit_text(ADMIN_HELP, reply_markup=admin_kb())
    await cb.answer()


@router.callback_query(F.data.startswith("a:obj:"))
async def a_obj(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    obj_id = int(cb.data.split(":")[2])
    with db() as c:
        w = c.execute(
            "SELECT COUNT(*) n, IFNULL(SUM(hours),0) h, IFNULL(SUM(expenses),0) m "
            "FROM reports WHERE object_id=?", (obj_id,)).fetchone()
        e = c.execute(
            "SELECT IFNULL(SUM(amount),0) m FROM expenses WHERE object_id=?", (obj_id,)).fetchone()
    kb = InlineKeyboardBuilder()
    kb.button(text="✏️ Переименовать", callback_data=f"a:ren:{obj_id}")
    kb.button(text="📦 В архив", callback_data=f"a:arc:{obj_id}")
    kb.button(text="⬅️ К списку", callback_data="a:list")
    kb.adjust(1)
    await cb.message.edit_text(
        f"🏗 {object_name(obj_id)}\n\n"
        f"Записей о работе: {w['n']}\n"
        f"Всего часов: {w['h']:g}\n"
        f"Всего расходов: {w['m'] + e['m']:g}",
        reply_markup=kb.as_markup(),
    )
    await cb.answer()


@router.callback_query(F.data.startswith("a:ren:"))
async def a_rename(cb: CallbackQuery, state: FSMContext):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    obj_id = int(cb.data.split(":")[2])
    await state.set_state(Adm.rename)
    await state.update_data(obj_id=obj_id)
    await cb.message.edit_text(f"Сейчас: {object_name(obj_id)}\n\nНапишите новое название:")
    await cb.answer()


@router.message(Adm.rename, F.text)
async def a_rename_save(message: Message, state: FSMContext):
    data = await state.get_data()
    name = " ".join((message.text or "").split())
    if len(name) < 2:
        await message.answer("Слишком короткое название, напишите ещё раз:")
        return
    with db() as c:
        c.execute("UPDATE objects SET name=? WHERE id=?", (name, data["obj_id"]))
    await state.clear()
    await message.answer(f"Переименован: {name}")
    await back_to_admin(message)


@router.callback_query(F.data.startswith("a:arc:"))
async def a_archive(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    obj_id = int(cb.data.split(":")[2])
    with db() as c:
        c.execute("UPDATE objects SET active=0 WHERE id=?", (obj_id,))
    await cb.message.edit_text(
        f"Объект «{object_name(obj_id)}» убран из списка. "
        "Монтажники его больше не увидят, записи сохранены — вернуть можно из архива.",
        reply_markup=admin_kb(),
    )
    await cb.answer()


@router.callback_query(F.data == "a:arch")
async def a_arch_list(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    with db() as c:
        objs = c.execute("SELECT id, name FROM objects WHERE active=0 ORDER BY name").fetchall()
    if not objs:
        await cb.message.edit_text("Архив пуст.", reply_markup=admin_kb())
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    for o in objs:
        kb.button(text=f"♻️ {o['name']}", callback_data=f"a:res:{o['id']}")
    kb.button(text="⬅️ Назад", callback_data="a:back")
    kb.adjust(1)
    await cb.message.edit_text("Нажмите, чтобы вернуть объект в работу:", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.startswith("a:res:"))
async def a_restore(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    obj_id = int(cb.data.split(":")[2])
    with db() as c:
        c.execute("UPDATE objects SET active=1 WHERE id=?", (obj_id,))
    await cb.message.edit_text(f"Объект «{object_name(obj_id)}» снова в работе.",
                               reply_markup=admin_kb())
    await cb.answer()


@router.callback_query(F.data == "a:people")
async def a_people(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    with db() as c:
        rows = c.execute(
            """SELECT u.fio, u.status, IFNULL(SUM(r.hours),0) h, COUNT(r.id) n
               FROM users u LEFT JOIN reports r ON r.user_id = u.user_id
               GROUP BY u.user_id ORDER BY u.fio"""
        ).fetchall()
    if not rows:
        await cb.message.edit_text("Пока никто не зарегистрировался.", reply_markup=admin_kb())
        await cb.answer()
        return
    marks = {"approved": "", "pending": " ⏳ ждёт подтверждения", "rejected": " ⛔️ отклонён"}
    text = "👥 Монтажники:\n\n" + "\n".join(
        f"• {r['fio']} — записей {r['n']}, часов {r['h']:g}{marks.get(r['status'], '')}"
        for r in rows
    )
    await cb.message.edit_text(text, reply_markup=admin_kb())
    await cb.answer()


# --- заявки на доступ ---
@router.callback_query(F.data.startswith("u:"))
async def decide_user(cb: CallbackQuery, bot: Bot):
    if not is_admin(cb.from_user.id):
        await cb.answer("Недоступно", show_alert=True)
        return
    _, action, uid = cb.data.split(":")
    uid = int(uid)
    fio = get_fio(uid) or f"ID {uid}"
    if user_status(uid) != "pending":
        await cb.message.edit_text(f"{fio} — заявка уже обработана.")
        await cb.answer()
        return
    who = get_fio(cb.from_user.id) or "руководитель"
    if action == "ok":
        set_status(uid, "approved")
        await cb.message.edit_text(f"✅ {fio} допущен к работе (принял: {who})")
        try:
            await bot.send_message(
                uid, "Доступ открыт. Нажмите /start, чтобы внести первый отчёт."
            )
        except Exception:
            pass
    else:
        set_status(uid, "rejected")
        await cb.message.edit_text(f"⛔️ {fio} отклонён (решение: {who})")
        try:
            await bot.send_message(uid, DENIED_TEXT)
        except Exception:
            pass
    await cb.answer()
    # остальным начальникам сообщаем, что решение принято
    for admin in admin_ids():
        if admin != cb.from_user.id:
            try:
                mark = "✅ принял" if action == "ok" else "⛔️ отклонил"
                await bot.send_message(admin, f"{who} {mark} заявку: {fio}")
            except Exception:
                pass


@router.callback_query(F.data == "a:req")
async def a_requests(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    rows = pending_users()
    if not rows:
        await cb.message.edit_text("Новых заявок нет.", reply_markup=admin_kb())
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    for r in rows:
        kb.button(text=f"✅ {r['fio']}", callback_data=f"u:ok:{r['user_id']}")
        kb.button(text=f"⛔️ {r['fio']}", callback_data=f"u:no:{r['user_id']}")
    kb.button(text="⬅️ Назад", callback_data="a:back")
    kb.adjust(2)
    await cb.message.edit_text(
        "🔔 Ждут подтверждения:\n\n" + "\n".join(f"• {r['fio']}" for r in rows),
        reply_markup=kb.as_markup(),
    )
    await cb.answer()


# --- начальники ---
def person_label(user_id):
    return get_fio(user_id) or f"ID {user_id}"


@router.callback_query(F.data == "a:bosses")
async def a_bosses(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    with db() as c:
        extra = c.execute("SELECT user_id, fio FROM admins ORDER BY fio").fetchall()
    lines = ["👑 Доступ к панели есть у:", ""]
    for uid in sorted(ADMIN_IDS):
        lines.append(f"• {person_label(uid)} — главный, снять нельзя")
    for r in extra:
        lines.append(f"• {r['fio'] or person_label(r['user_id'])}")
    kb = InlineKeyboardBuilder()
    if is_owner(cb.from_user.id):
        kb.button(text="➕ Назначить начальника", callback_data="a:promote")
        for r in extra:
            kb.button(text=f"➖ Снять: {r['fio'] or r['user_id']}",
                      callback_data=f"a:demote:{r['user_id']}")
    else:
        lines += ["", "Назначать и снимать может только главный администратор."]
    kb.button(text="⬅️ Назад", callback_data="a:back")
    kb.adjust(1)
    await cb.message.edit_text("\n".join(lines), reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data == "a:promote")
async def a_promote_list(cb: CallbackQuery):
    if not is_owner(cb.from_user.id):
        await cb.answer("Только главный администратор", show_alert=True)
        return
    with db() as c:
        rows = c.execute("SELECT user_id, fio FROM users ORDER BY fio").fetchall()
    candidates = [r for r in rows if not is_admin(r["user_id"])]
    if not candidates:
        await cb.message.edit_text(
            "Некого назначить.\n\n"
            "Пусть коллега сначала откроет бот, нажмёт /start и впишет своё ФИО — "
            "после этого он появится в этом списке.",
            reply_markup=admin_kb(),
        )
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    for r in candidates:
        kb.button(text=r["fio"], callback_data=f"a:prom:{r['user_id']}")
    kb.button(text="⬅️ Назад", callback_data="a:bosses")
    kb.adjust(1)
    await cb.message.edit_text(
        "Кому дать доступ к панели администратора?\n"
        "У этого человека появятся объекты, отчёты и уведомления о новых записях.",
        reply_markup=kb.as_markup(),
    )
    await cb.answer()


@router.callback_query(F.data.startswith("a:prom:"))
async def a_promote(cb: CallbackQuery, bot: Bot):
    if not is_owner(cb.from_user.id):
        await cb.answer("Только главный администратор", show_alert=True)
        return
    uid = int(cb.data.split(":")[2])
    fio = person_label(uid)
    with db() as c:
        c.execute(
            "INSERT OR REPLACE INTO admins (user_id, fio, added_by, created_at) VALUES (?,?,?,?)",
            (uid, fio, person_label(cb.from_user.id), datetime.now().strftime("%Y-%m-%d %H:%M")),
        )
    await apply_commands(bot, uid)
    await cb.message.edit_text(f"{fio} теперь начальник — панель /admin ему доступна.",
                               reply_markup=admin_kb())
    await cb.answer()
    try:
        await bot.send_message(
            uid, "Вам открыт доступ к панели администратора. Откройте её командой /admin"
        )
    except Exception:
        pass


@router.callback_query(F.data.startswith("a:demote:"))
async def a_demote(cb: CallbackQuery, bot: Bot):
    if not is_owner(cb.from_user.id):
        await cb.answer("Только главный администратор", show_alert=True)
        return
    uid = int(cb.data.split(":")[2])
    fio = person_label(uid)
    with db() as c:
        c.execute("DELETE FROM admins WHERE user_id=?", (uid,))
    await apply_commands(bot, uid)
    await cb.message.edit_text(f"{fio} больше не начальник — панель ему недоступна.",
                               reply_markup=admin_kb())
    await cb.answer()


def add_sheet(wb, title, headers, rows, totals=None, widths=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    if totals:
        ws.append([])
        ws.append(totals)
        ws[ws.max_row][0].font = Font(bold=True)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    return ws


@router.message(Command("report"))
async def report_cmd(message: Message):
    if not is_admin(message.from_user.id):
        return
    await send_report(message, (message.text or "").partition(" ")[2].strip())


@router.callback_query(F.data == "a:rep")
async def a_report_menu(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    kb.button(text="📊 Всё целиком", callback_data="rep:all")
    kb.button(text="📅 За текущий месяц", callback_data="rep:month")
    kb.button(text="🏗 По объекту", callback_data="rep:obj")
    kb.button(text="👤 По монтажнику", callback_data="rep:user")
    kb.button(text="⬅️ Назад", callback_data="a:back")
    kb.adjust(1)
    await cb.message.edit_text(
        "Какой отчёт выгрузить?\n\n"
        "За другой месяц — команда /report 09.2026",
        reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.in_({"rep:all", "rep:month"}))
async def a_report_simple(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    await cb.answer("Готовлю файл…")
    arg = date.today().strftime("%m.%Y") if cb.data == "rep:month" else ""
    await send_report(cb.message, arg)
    await cb.message.answer("Что дальше?", reply_markup=admin_kb())


@router.callback_query(F.data == "rep:obj")
async def a_report_pick_obj(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    with db() as c:
        objs = c.execute("SELECT id, name FROM objects ORDER BY active DESC, name").fetchall()
    if not objs:
        await cb.answer("Объектов нет", show_alert=True)
        return
    kb = InlineKeyboardBuilder()
    for o in objs:
        kb.button(text=o["name"], callback_data=f"repo:{o['id']}")
    kb.button(text="⬅️ Назад", callback_data="a:rep")
    kb.adjust(1)
    await cb.message.edit_text("По какому объекту сделать отчёт?", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data == "rep:user")
async def a_report_pick_user(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    people = approved_users()
    if not people:
        await cb.answer("Монтажников нет", show_alert=True)
        return
    kb = InlineKeyboardBuilder()
    for r in people:
        kb.button(text=r["fio"], callback_data=f"repu:{r['user_id']}")
    kb.button(text="⬅️ Назад", callback_data="a:rep")
    kb.adjust(1)
    await cb.message.edit_text("По кому сделать отчёт?", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.startswith("repo:"))
async def a_report_by_obj(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    await cb.answer("Готовлю файл…")
    await send_report(cb.message, "", obj_id=int(cb.data.split(":")[1]))
    await cb.message.answer("Что дальше?", reply_markup=admin_kb())


@router.callback_query(F.data.startswith("repu:"))
async def a_report_by_user(cb: CallbackQuery):
    if not is_admin(cb.from_user.id):
        await cb.answer()
        return
    await cb.answer("Готовлю файл…")
    await send_report(cb.message, "", user_id=int(cb.data.split(":")[1]))
    await cb.message.answer("Что дальше?", reply_markup=admin_kb())


def fetch_rows(period="", obj_id=None, user_id=None):
    """Строки для отчёта с нужными фильтрами."""
    where_w, where_e, params = [], [], []
    if period:
        where_w.append("substr(r.work_date,1,7)=?")
        where_e.append("substr(e.exp_date,1,7)=?")
        params.append(period)
    if obj_id is not None:
        where_w.append("r.object_id=?")
        where_e.append("e.object_id=?")
        params.append(obj_id)
    if user_id is not None:
        where_w.append("r.user_id=?")
        where_e.append("e.user_id=?")
        params.append(user_id)
    cond_w = (" WHERE " + " AND ".join(where_w)) if where_w else ""
    cond_e = (" WHERE " + " AND ".join(where_e)) if where_e else ""

    q1 = """SELECT r.*, o.name AS obj_name FROM reports r
            LEFT JOIN objects o ON o.id = r.object_id"""
    q2 = """SELECT e.*, o.name AS obj_name FROM expenses e
            LEFT JOIN objects o ON o.id = e.object_id"""
    with db() as c:
        works = c.execute(q1 + cond_w + " ORDER BY r.work_date", tuple(params)).fetchall()
        exps = c.execute(q2 + cond_e + " ORDER BY e.exp_date", tuple(params)).fetchall()
    return works, exps


async def send_report(message: Message, arg: str, obj_id=None, user_id=None):
    period = ""
    if arg:
        try:
            period = datetime.strptime(arg, "%m.%Y").strftime("%Y-%m")
        except ValueError:
            await message.answer("Формат: /report 09.2026")
            return
    works, exps = fetch_rows(period, obj_id, user_id)

    scope = []
    if obj_id is not None:
        scope.append(f"объект: {object_name(obj_id)}")
    if user_id is not None:
        scope.append(f"монтажник: {get_fio(user_id) or user_id}")
    if period:
        scope.append(f"месяц: {arg}")
    scope_text = ("\n" + ", ".join(scope).capitalize()) if scope else ""

    if not works and not exps:
        await message.answer("Записей нет." + scope_text)
        return

    buf = build_workbook(works, exps)
    parts = ["otchet"]
    if obj_id is not None:
        parts.append(f"obj{obj_id}")
    if user_id is not None:
        parts.append(f"user{user_id}")
    parts.append(arg.replace(".", "_") or "vse")
    total_money = sum(e["amount"] for e in exps) + sum(r["expenses"] for r in works)
    await message.answer_document(
        BufferedInputFile(buf.getvalue(), filename="_".join(parts) + ".xlsx"),
        caption=(f"Работ: {len(works)} (часов: {sum(r['hours'] for r in works):g})\n"
                 f"Расходов: {len(exps)}\n"
                 f"Всего потрачено: {total_money:g}" + scope_text),
    )


def build_workbook(works, exps):
    """Собирает файл Excel: листы «Работы», «Сводка», «Расходы»."""
    # сводка: по каждому монтажнику и объекту — часы и деньги
    summary = {}
    for r in works:
        key = (r["user_name"], r["obj_name"])
        s = summary.setdefault(key, {"hours": 0.0, "money": 0.0, "days": set()})
        s["hours"] += r["hours"] or 0
        s["money"] += r["expenses"] or 0
        s["days"].add(r["work_date"])
    for e in exps:
        key = (e["user_name"], e["obj_name"])
        s = summary.setdefault(key, {"hours": 0.0, "money": 0.0, "days": set()})
        s["money"] += e["amount"] or 0
    # группируем по объекту, после каждого объекта — строка «Итого по объекту»
    summary_rows = []
    bold_rows = []
    by_object = {}
    for (fio, obj), v in summary.items():
        by_object.setdefault(obj, []).append((fio, v))
    for obj in sorted(by_object, key=lambda x: (x is None, str(x))):
        people = sorted(by_object[obj], key=lambda p: str(p[0]))
        for fio, v in people:
            summary_rows.append([obj, fio, len(v["days"]), round(v["hours"], 2), v["money"]])
        summary_rows.append([
            f"ИТОГО по объекту: {obj}", "",
            len({d for _, v in people for d in v["days"]}),
            round(sum(v["hours"] for _, v in people), 2),
            sum(v["money"] for _, v in people),
        ])
        bold_rows.append(len(summary_rows) + 1)   # +1 — строка заголовков
        summary_rows.append([])

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(
        wb, "Работы",
        ["Дата", "Объект", "Монтажник", "С", "По", "Часов", "Расходы", "На что расход",
         "Описание и с кем работал", "Внесено", "Внёс за него"],
        [[fmt_date(r["work_date"]), r["obj_name"], r["user_name"], r["time_from"], r["time_to"],
          r["hours"], r["expenses"], r["exp_note"] or "", r["description"], r["created_at"],
          r["entered_by"] or ""] for r in works],
        ["ИТОГО", "", "", "", "", sum(r["hours"] for r in works),
         sum(r["expenses"] for r in works)],
        (12, 26, 24, 7, 7, 8, 10, 32, 42, 17, 22),
    )
    ws_sum = add_sheet(
        wb, "Сводка",
        ["Объект", "Монтажник", "Дней", "Часов", "Потрачено"],
        summary_rows,
        None,
        (34, 26, 8, 10, 12),
    )
    for row_idx in bold_rows:
        for cell in ws_sum[row_idx]:
            cell.font = Font(bold=True)
    add_sheet(
        wb, "Расходы",
        ["Дата", "Объект", "Монтажник", "Сумма", "На что", "Внесено", "Внёс за него"],
        [[fmt_date(e["exp_date"]), e["obj_name"], e["user_name"], e["amount"],
          e["comment"], e["created_at"], e["entered_by"] or ""] for e in exps],
        ["ИТОГО", "", "", sum(e["amount"] for e in exps)],
        (12, 26, 24, 11, 45, 17, 22),
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf


# ---------- Яндекс.Диск ----------
# Бот держит на Диске один всегда актуальный файл отчёта и копию базы.
YADISK_API = "https://cloud-api.yandex.net/v1/disk/resources"
_need_sync = False          # появились новые записи — файл пора обновить
_last_backup = ""           # дата последней копии базы


def mark_changed():
    """Отмечает, что данные изменились и файл на Диске устарел."""
    global _need_sync
    _need_sync = True


async def yadisk_put(session, disk_path, data: bytes):
    """Кладёт файл на Диск, перезаписывая прежний."""
    headers = {"Authorization": f"OAuth {YADISK_TOKEN}"}
    # папка может ещё не существовать — 409 означает, что она уже есть
    folder = disk_path.rsplit("/", 1)[0]
    if folder:
        async with session.put(YADISK_API, params={"path": folder}, headers=headers):
            pass
    async with session.get(f"{YADISK_API}/upload", headers=headers,
                           params={"path": disk_path, "overwrite": "true"}) as resp:
        if resp.status != 200:
            logging.warning("Яндекс.Диск: не дал ссылку на загрузку (%s)", resp.status)
            return False
        href = (await resp.json())["href"]
    async with session.put(href, data=data) as resp:
        ok = resp.status in (201, 202)
        if not ok:
            logging.warning("Яндекс.Диск: файл не загрузился (%s)", resp.status)
        return ok


async def sync_to_yadisk(force=False):
    """Обновляет на Диске файл отчёта, раз в сутки — копию базы."""
    global _need_sync, _last_backup
    if not YADISK_TOKEN or (not _need_sync and not force):
        return
    _need_sync = False
    works, exps = fetch_rows()
    if not works and not exps:
        return
    import aiohttp
    try:
        async with aiohttp.ClientSession() as session:
            buf = build_workbook(works, exps)
            await yadisk_put(session, f"{YADISK_FOLDER}/{YADISK_FILE}", buf.getvalue())
            today = date.today().isoformat()
            if _last_backup != today and os.path.exists(DB_PATH):
                with open(DB_PATH, "rb") as f:
                    if await yadisk_put(session, f"{YADISK_FOLDER}/Копия базы/{today}.db", f.read()):
                        _last_backup = today
    except Exception as e:
        logging.warning("Яндекс.Диск: синхронизация не удалась: %s", e)


async def yadisk_loop():
    """Фоновая задача: проверяет изменения раз в минуту."""
    while True:
        await asyncio.sleep(SYNC_SECONDS)
        await sync_to_yadisk()


@router.message(Command("disk"))
async def disk_now(message: Message):
    """Обновить файл на Диске прямо сейчас."""
    if not is_admin(message.from_user.id):
        return
    if not YADISK_TOKEN:
        await message.answer(
            "Выгрузка на Яндекс.Диск не настроена.\n"
            "Добавьте на хостинге переменную YADISK_TOKEN с токеном Яндекс.Диска."
        )
        return
    await sync_to_yadisk(force=True)
    await message.answer(
        f"Файл на Яндекс.Диске обновлён:\n{YADISK_FOLDER}/{YADISK_FILE}"
    )


# ---------- Запуск ----------
async def main():
    logging.basicConfig(level=logging.INFO)
    if not BOT_TOKEN:
        raise SystemExit("Не задан токен: укажите переменную окружения BOT_TOKEN")
    init_db()
    bot = Bot(BOT_TOKEN)
    # обычный список команд — его видят монтажники
    await bot.set_my_commands(USER_COMMANDS)
    # начальникам дополнительно показываем /admin
    for uid in admin_ids():
        await apply_commands(bot, uid)
    dp = Dispatcher(storage=MemoryStorage())
    dp.include_router(router)
    if YADISK_TOKEN:
        logging.info("Яндекс.Диск подключён: %s/%s", YADISK_FOLDER, YADISK_FILE)
        asyncio.create_task(yadisk_loop())
        await sync_to_yadisk(force=True)
    else:
        logging.info("Яндекс.Диск не настроен (нет YADISK_TOKEN) — бот работает без выгрузки")
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
